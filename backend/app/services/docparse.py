"""Text extraction from uploaded documents: PDF, Word, PowerPoint, plain text, Markdown and HTML.

`extract_text(data, filename)` returns (text, title). It raises `DocParseError` with a message the
user can act on when a file can't be read.
"""

import io
import re
from pathlib import PurePath

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
SUPPORTED = {".pdf", ".docx", ".pptx", ".txt", ".md", ".markdown", ".html", ".htm", ".csv"}


class DocParseError(Exception):
    pass


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode(data: bytes) -> str:
    encodings = ("utf-16", "utf-8-sig") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1252")
    for enc in encodings:
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _pdf(data: bytes) -> tuple[str, str | None]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise DocParseError("This PDF is password-protected. Remove the password and upload it again.")
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                continue
        title = None
        try:
            title = (reader.metadata.title if reader.metadata else None) or None
        except Exception:
            title = None
        return "\n\n".join(pages), (str(title).strip() or None) if title else None
    except DocParseError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as e:
        raise DocParseError(f"Couldn't read this PDF ({e}).")


def _docx(data: bytes) -> tuple[str, str | None]:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise DocParseError(f"Couldn't read this Word file ({e}). Save it as .docx and try again.")
    parts: list[str] = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        parts.append(f"\n{t}\n" if style.startswith("heading") or style == "title" else t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    title = (d.core_properties.title or "").strip() or None
    return "\n".join(parts), title


def _pptx(data: bytes) -> tuple[str, str | None]:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise DocParseError(f"Couldn't read this PowerPoint file ({e}). Save it as .pptx and try again.")
    slides: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        lines: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    lines.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    lines.append(f"Notes: {notes}")
        except Exception:
            pass
        if lines:
            slides.append(f"Slide {i}\n" + "\n".join(lines))
    title = (prs.core_properties.title or "").strip() or None
    return "\n\n".join(slides), title


def _html(data: bytes) -> tuple[str, str | None]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "lxml")
    title = soup.title.get_text(strip=True) if soup.title else None
    for tag in soup(["script", "style", "noscript", "svg", "template"]):
        tag.decompose()
    return soup.get_text("\n"), title or None


def extract_text(data: bytes, filename: str) -> tuple[str, str | None]:
    """Return (clean text, document title or None)."""
    if not data:
        raise DocParseError("The file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocParseError(f"The file is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB. Split it and upload the parts.")
    ext = PurePath(filename or "").suffix.lower()
    if ext not in SUPPORTED:
        raise DocParseError("Upload a PDF, Word (.docx), PowerPoint (.pptx), text, Markdown or HTML file.")
    if ext == ".pdf":
        text, title = _pdf(data)
    elif ext == ".docx":
        text, title = _docx(data)
    elif ext == ".pptx":
        text, title = _pptx(data)
    elif ext in (".html", ".htm"):
        text, title = _html(data)
    else:
        text, title = _decode(data), None
    text = _clean(text)
    if len(text) < 20:
        raise DocParseError(
            "No readable text was found. If this is a scanned document, export it with text (OCR) and upload it again."
        )
    return text, title
