"""Company Intelligence: crawler, extraction, knowledge retrieval, competitors, ideas, API permissions and tenancy.

The network and the AI model are mocked: the crawler talks to an httpx.MockTransport serving a small fake
site, DNS is stubbed to a public address, web search and PageSpeed are replaced, and the LLM router's
`complete` returns canned JSON chosen by the system prompt.
"""

import io
import json
import re
import uuid
from datetime import UTC, datetime

import httpx
import pytest
from sqlalchemy import select

import app.modules.intelligence as intel
from app.core.db import SessionLocal, set_tenant
from app.models import CompanyEntity, CompanyProfile, Competitor, ContentPillar, KnowledgeDocument, KnowledgeSource, Opportunity
from app.services import crawler, docparse, knowledge, search
from app.services.jobs import run_inline
from app.services.llm import router as llm_router
from app.services.llm.router import LlmResult, NoProviderError
from tests.conftest import auth, login

ADMIN_PW = "Str0ngPassw0rd!"
SITE = "https://acme-hr.test"
TODAY = datetime.now(UTC).date().isoformat()

LONG = (
    "Acme HRMS handles payroll compliance, attendance, leave and statutory filings for Indian companies. "
    "Managers approve requests on mobile and HR teams close payroll in hours instead of days. "
)


def page(title: str, body: str, *, desc: str = "", extra_head: str = "", links: str = "") -> str:
    return (f"<html lang='en'><head><title>{title}</title><meta name='description' content='{desc}'>{extra_head}</head>"
            f"<body><nav>{links}</nav><main><h1>{title}</h1><p>{body}</p></main></body></html>")


NAV = ("<a href='/about'>About</a><a href='/products/hrms'>HRMS</a><a href='/clients'>Clients</a>"
       "<a href='/blog/post-1'>Blog</a><a href='/missing'>Old page</a><a href='/brochure.pdf'>Brochure</a>"
       "<a href='https://www.linkedin.com/company/acme-hr'>LinkedIn</a>"
       "<a href='https://www.facebook.com/sharer/sharer.php?u=x'>Share</a>")

SITE_PAGES = {
    "/": page("Acme HR – AI HRMS for growing teams", LONG * 3 + "<img src='/a.png'>", desc="AI HRMS and payroll for India",
              links=NAV, extra_head="<script type='application/ld+json'>{\"@type\":\"Organization\",\"sameAs\":"
                                    "[\"https://www.facebook.com/acmehr\"]}</script>"),
    "/about": page("About Acme HR", "Acme HR is a Bengaluru company building AI tools for HR teams. " * 6, links=NAV),
    "/products/hrms": page("Acme HRMS", LONG * 4, links=NAV),
    "/clients": page("Our clients", "Our clients include Globex Corporation and Initech, who run payroll on Acme. " * 5, links=NAV),
    "/blog/post-1": page("Five payroll mistakes", "Payroll mistakes cost money. Here is how to avoid them. " * 8,
                         extra_head=f"<meta property='article:published_time' content='{TODAY}T09:00:00Z'>", links=NAV),
    "/private/secret": page("Secret", "should never be fetched " * 20),
}
ROBOTS = "User-agent: *\nDisallow: /private/\nSitemap: https://acme-hr.test/sitemap.xml\n"
SITEMAP = ("<?xml version='1.0'?><urlset xmlns='http://www.sitemaps.org/schemas/sitemap/0.9'>"
           "<url><loc>https://acme-hr.test/about</loc></url><url><loc>https://acme-hr.test/products/hrms</loc></url>"
           "<url><loc>https://acme-hr.test/private/secret</loc></url></urlset>")

RIVALS = {
    "rival1.test": page("Rival One HRMS", "Rival One sells HRMS and payroll software to Indian SMEs. " * 6),
    "rival2.test": page("Rival Two People Platform", "Rival Two is an AI people platform with payroll and attendance. " * 6,
                        links="<a href='/about-us'>About</a>"),
    "news-daily.test": page("Daily HR news", "Headlines about HR. " * 20),
}

REQUESTED: list[str] = []


def _site_handler(request: httpx.Request) -> httpx.Response:
    url = request.url
    REQUESTED.append(str(url))
    host, path = url.host.removeprefix("www."), url.path
    if host == "acme-hr.test":
        if url.scheme == "http":
            return httpx.Response(301, headers={"location": f"{SITE}{path}"})
        if path == "/robots.txt":
            return httpx.Response(200, text=ROBOTS, headers={"content-type": "text/plain"})
        if path == "/sitemap.xml":
            return httpx.Response(200, text=SITEMAP, headers={"content-type": "application/xml"})
        if path in SITE_PAGES:
            return httpx.Response(200, text=SITE_PAGES[path], headers={"content-type": "text/html; charset=utf-8"})
        return httpx.Response(404, text="not found", headers={"content-type": "text/html"})
    if host in RIVALS:
        if path == "/robots.txt":
            return httpx.Response(404)
        if path == "/":
            return httpx.Response(200, text=RIVALS[host], headers={"content-type": "text/html"})
        if path == "/about-us":
            return httpx.Response(200, text=page("About Rival Two", "We serve HR teams across India. " * 10),
                                  headers={"content-type": "text/html"})
        if path == "/new-feature" and host == "rival1.test" and STATE.get("rival1_new_page"):
            return httpx.Response(200, text=page("Rival One AI copilot", "Our new AI copilot for HR. " * 10),
                                  headers={"content-type": "text/html"})
    return httpx.Response(404, headers={"content-type": "text/html"})


STATE: dict = {}
SEARCH_RESULTS = [
    {"title": "Acme HR", "url": "https://acme-hr.test/", "snippet": "us"},
    {"title": "Acme on LinkedIn", "url": "https://www.linkedin.com/company/acme-hr", "snippet": "social"},
    {"title": "Best HRMS - G2", "url": "https://www.g2.com/categories/hrms", "snippet": "review site"},
    {"title": "HRMS - Wikipedia", "url": "https://en.wikipedia.org/wiki/HRMS", "snippet": "encyclopedia"},
    {"title": "Rival One", "url": "https://rival1.test/", "snippet": "HRMS for SMEs"},
    {"title": "Rival Two", "url": "https://www.rival2.test/pricing", "snippet": "people platform"},
    {"title": "Rival Three", "url": "https://rival3.test/", "snippet": "rejected earlier"},
    {"title": "HR news", "url": "https://news-daily.test/story", "snippet": "news"},
]


def _fake_llm_payload(system: str, prompt: str) -> dict:
    STATE.setdefault("calls", []).append(system[:40])
    if system == intel.EXTRACT_SYSTEM:
        if "company's documents" in prompt:
            return {"summary": "", "industries": ["Retail"], "entities": [
                {"type": "client", "name": "Secret Corp", "description": "Pilot customer"},
                {"type": "product", "name": "Acme HRMS", "description": "HRMS"}]}
        return {
            "summary": "Acme HR builds an AI-powered HRMS with payroll for Indian companies.",
            "positioning": "The fastest way for Indian SMEs to run compliant payroll.",
            "industries": ["Manufacturing", "IT services"], "audiences": ["HR heads"], "differentiators": ["AI assistant"],
            "geography": "India", "entities": [
                {"type": "product", "name": "Acme HRMS", "description": "Payroll, attendance and leave", "url": f"{SITE}/products/hrms"},
                {"type": "studio", "name": "AI Studio", "description": "Custom AI agents"},
                {"type": "client", "name": "Globex Corporation", "description": "Manufacturer"},
                {"type": "client", "name": "Initech"},
                {"type": "use_case", "name": "Payroll compliance"},
                {"type": "nonsense", "name": "Ignored"},
                {"type": "industry", "name": ""},
            ]}
    if system == intel.MERGE_SYSTEM:
        return {"summary": "Acme HR builds an AI-powered HRMS with payroll for Indian companies.",
                "positioning": "Compliant payroll for SMEs.", "industries": ["Manufacturing", "IT services", "Retail"],
                "audiences": ["HR heads"], "differentiators": ["AI assistant"], "geography": "India"}
    if system == intel.QUERIES_SYSTEM:
        return {"queries": ["AI HRMS software India", "payroll software India", "HR analytics platform India"]}
    if system == intel.SCORE_SYSTEM:
        doms = re.findall(r"Candidate domain: (\S+)", prompt)
        return {"candidates": [{"domain": d, "name": d.split(".")[0].title(), "overlap": 0.85 if "rival" in d else 0.1,
                                "reason": "Sells HRMS to the same buyers", "offerings": ["HRMS"], "positioning": "HR suite"}
                               for d in doms]}
    if system == intel.COMPETITOR_SYSTEM:
        offerings = ["HRMS", "Payroll"] + (["AI copilot"] if STATE.get("rival1_new_page") else [])
        return {"offerings": offerings, "positioning": "HR suite for SMEs", "content_themes": ["Payroll compliance", "Remote work"],
                "blog_topics": ["Gratuity rules explained"], "keywords": ["hrms"], "pricing_signals": ["per employee per month"]}
    if system == intel.IDEAS_SYSTEM:
        return {"ideas": [
            {"title": "How Acme HRMS closes payroll in hours", "angle": "Speed", "source": "idea", "score": 0.9, "pillar": "Product",
             "keywords": ["payroll"], "formats": ["carousel"]},
            {"title": "How Acme HRMS closes payroll in HOURS!", "angle": "Duplicate", "source": "idea", "score": 0.8},
            {"title": "Gratuity rules explained for SMEs", "angle": "Gap", "source": "gap", "score": 0.7, "pillar": "Unknown pillar"},
            {"title": "An already existing manual idea", "angle": "dup of existing", "source": "idea", "score": 0.5},
            {"title": "Search-led idea", "source": "gsc", "score": "85"},
            {"title": "", "source": "idea"},
            "not a dict",
        ]}
    raise AssertionError(f"unexpected prompt: {system[:60]}")


async def fake_complete(session, task, prompt=None, *, messages=None, system=None, json_mode=False, temperature=0.7, max_tokens=1500):
    if STATE.get("no_ai"):
        raise NoProviderError("No AI text provider is set up.")
    payload = _fake_llm_payload(system or "", prompt or "")
    return LlmResult(json.dumps(payload), "fake", uuid.uuid4(), "fake-model")


async def fake_search(session, query, n=10):
    STATE.setdefault("queries", []).append(query)
    return SEARCH_RESULTS


async def fake_pagespeed(url, key):
    return {"performance": 62, "accessibility": 90, "best_practices": 95, "seo": 92, "lcp_ms": 3400, "cls": 0.02}


@pytest.fixture(autouse=True)
def mocks(monkeypatch):
    async def public_dns(host):
        return ["93.184.216.34"]

    async def no_render(self, url):
        return None

    monkeypatch.setattr(crawler, "TRANSPORT", httpx.MockTransport(_site_handler))
    monkeypatch.setattr(crawler, "resolve_host", public_dns)
    monkeypatch.setattr(crawler._Renderer, "render", no_render)
    monkeypatch.setattr(llm_router, "complete", fake_complete)
    monkeypatch.setattr(search, "web_search", fake_search)
    monkeypatch.setattr(intel, "fetch_pagespeed", fake_pagespeed)
    STATE.pop("no_ai", None)
    yield


@pytest.fixture
async def hrp(client, admin):
    return await login(client, "admin@hrpai.in", ADMIN_PW)


def org_of(tokens) -> uuid.UUID:
    return uuid.UUID(tokens["organization"]["id"])


async def tenant_session(org_id):
    s = SessionLocal()
    set_tenant(s, org_id)
    return s


# ------------------------------------------------------------------ pure helpers
async def test_chunking_splits_on_paragraphs_with_overlap():
    paras = [f"Paragraph {i}. " + ("word " * 60).strip() for i in range(12)]
    chunks = knowledge.chunk_text("\n\n".join(paras))
    assert len(chunks) > 3
    assert all(len(c) <= knowledge.CHUNK_SIZE + 10 for c in chunks)
    # Each chunk after the first starts with the tail of the previous one.
    for prev, cur in zip(chunks, chunks[1:], strict=False):
        assert cur.split("\n\n")[0] in prev
    # A single giant paragraph is still split, and nothing is lost.
    giant = "Sentence number one is here. " * 200
    parts = knowledge.chunk_text(giant)
    assert len(parts) > 5 and all(len(p) <= knowledge.CHUNK_SIZE + 10 for p in parts)
    assert knowledge.chunk_text("   ") == []


async def test_url_helpers_and_ssrf_rules():
    assert crawler.registrable_domain("blog.www.hrpai.co.in") == "hrpai.co.in"
    assert crawler.registrable_domain("www.example.com") == "example.com"
    assert crawler.normalize_url("HTTPS://Example.com:443/a//b?utm_source=x&id=2#top") == "https://example.com/a/b?id=2"
    assert crawler.social_platform("https://www.linkedin.com/company/acme") == "linkedin"
    assert crawler.social_platform("https://www.facebook.com/sharer/sharer.php?u=1") is None
    assert crawler.has_skipped_extension("https://x.com/file.PDF")
    for bad in ("http://127.0.0.1/", "http://localhost:8000", "ftp://example.com", "http://10.0.0.5", "http://[::1]/",
                "http://169.254.169.254/latest", "http://user:pw@example.com"):
        with pytest.raises(crawler.UnsafeURLError):
            await crawler.check_public_url(bad)
    assert await crawler.check_public_url("acme-hr.test") == "https://acme-hr.test/"


async def test_docparse_formats():
    import docx
    from pptx import Presentation

    d = docx.Document()
    d.add_heading("Acme brochure", 0)
    d.add_paragraph("Acme HRMS runs payroll for 500 companies across India.")
    buf = io.BytesIO()
    d.save(buf)
    text, _ = docparse.extract_text(buf.getvalue(), "brochure.docx")
    assert "payroll for 500 companies" in text

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "AI Studio"
    slide.placeholders[1].text = "Custom agents for HR workflows and approvals"
    buf = io.BytesIO()
    prs.save(buf)
    text, _ = docparse.extract_text(buf.getvalue(), "deck.pptx")
    assert "Slide 1" in text and "Custom agents" in text

    text, title = docparse.extract_text(b"<html><title>T</title><script>x()</script><p>Hello there, general text.</p></html>", "a.html")
    assert "Hello there" in text and "x()" not in text and title == "T"
    with pytest.raises(docparse.DocParseError):
        docparse.extract_text(b"\x00\x01", "file.exe")
    with pytest.raises(docparse.DocParseError):
        docparse.extract_text(b"%PDF-1.4 broken", "broken.pdf")


# ------------------------------------------------------------------ crawl → extract → context → retrieve
async def test_crawl_creates_documents(client, hrp):
    org = org_of(hrp)
    r = await client.post("/intel/sources/website", headers=auth(hrp), json={"url": "http://acme-hr.test", "max_pages": 20})
    assert r.status_code == 202, r.text
    sid = r.json()["source"]["id"]
    REQUESTED.clear()
    stats = await run_inline("intel.crawl", org, {"source_id": sid, "chain": False})
    assert stats["pages"] == 5 and stats["new"] == 5

    async with await tenant_session(org) as s:
        src = await s.get(KnowledgeSource, uuid.UUID(sid))
        assert src.status == "ready" and src.page_count == 5 and src.last_crawled_at
        crawl = src.meta["crawl"]
        assert crawl["robots_found"] and crawl["sitemap_found"] and crawl["https"]
        assert any(b["url"].endswith("/missing") for b in crawl["broken_links"])
        assert src.meta["social_links"]["linkedin"] == "https://www.linkedin.com/company/acme-hr"
        assert src.meta["social_links"]["facebook"] == "https://www.facebook.com/acmehr"
        docs = (await s.scalars(select(KnowledgeDocument).where(KnowledgeDocument.source_id == src.id))).all()
        urls = {d.url for d in docs}
        assert f"{SITE}/products/hrms" in urls and f"{SITE}/private/secret" not in urls
        home = next(d for d in docs if d.url == f"{SITE}/")
        assert home.meta["images_missing_alt"] == 1 and "Organization" in home.meta["jsonld_types"]
        blog = next(d for d in docs if d.url.endswith("/blog/post-1"))
        assert blog.meta["dates"] == [TODAY]
    # robots.txt respected and binaries never fetched
    assert not any("/private/" in u for u in REQUESTED)
    assert not any(u.endswith(".pdf") for u in REQUESTED)

    # Re-crawl of an unchanged site keeps documents and reports nothing new.
    again = await run_inline("intel.crawl", org, {"source_id": sid, "chain": False})
    assert again["pages"] == 5 and again["new"] == 0 and again["changed"] == 0


async def test_upload_document(client, hrp):
    files = {"file": ("notes.txt", b"Internal notes. Secret Corp is piloting Acme HRMS for payroll in Pune.\n\nMore text here.", "text/plain")}
    r = await client.post("/intel/sources/upload", headers=auth(hrp), files=files)
    assert r.status_code == 201, r.text
    assert r.json()["kind"] == "upload" and r.json()["filename"] == "notes.txt"
    bad = await client.post("/intel/sources/upload", headers=auth(hrp), files={"file": ("x.exe", b"MZ....", "application/octet-stream")})
    assert bad.status_code == 400


async def test_extract_builds_profile_and_entities(client, hrp):
    org = org_of(hrp)
    result = await run_inline("intel.extract", org, {"chain": False})
    assert result["added"] >= 5
    async with await tenant_session(org) as s:
        profile = await s.scalar(select(CompanyProfile))
        assert profile.summary.startswith("Acme HR builds") and profile.geography == "India"
        assert "Retail" in profile.industries
        assert profile.social_links["linkedin"].endswith("/company/acme-hr")
        ents = {(e.type, e.name): e for e in (await s.scalars(select(CompanyEntity))).all()}
        assert ("nonsense", "Ignored") not in ents
        assert ents[("client", "Globex Corporation")].public_mention_ok is True  # D20: named on the website
        assert ents[("client", "Initech")].source == "website"
        assert ents[("client", "Secret Corp")].public_mention_ok is False  # only in an uploaded document
        assert ents[("product", "Acme HRMS")].url == f"{SITE}/products/hrms"
        assert len([k for k in ents if k == ("product", "Acme HRMS")]) == 1
        studio_id = ents[("studio", "AI Studio")].id

    # A person's edit survives the next extraction.
    r = await client.patch(f"/intel/entities/{studio_id}", headers=auth(hrp), json={"description": "Edited by a person"})
    assert r.status_code == 200 and r.json()["edited"] is True
    r = await client.patch("/intel/profile", headers=auth(hrp), json={"geography": "India and the Middle East"})
    assert r.json()["profile"]["edited_fields"] == ["geography"]
    await run_inline("intel.extract", org, {"chain": False})
    async with await tenant_session(org) as s:
        studio = await s.get(CompanyEntity, uuid.UUID(str(studio_id)))
        assert studio.description == "Edited by a person"
        assert (await s.scalar(select(CompanyProfile))).geography == "India and the Middle East"


async def test_company_context_excludes_private_clients(client, hrp):
    org = org_of(hrp)
    async with await tenant_session(org) as s:
        ctx = await knowledge.company_context(s)
        assert "Globex Corporation" in ctx and "Initech" in ctx
        assert "Secret Corp" not in ctx
        assert "Acme HRMS" in ctx and "Summary:" in ctx
        assert len(await knowledge.company_context(s, max_chars=300)) <= 300

    # Once anything is confirmed, only confirmed facts are used.
    ents = (await client.get("/intel/entities", headers=auth(hrp), params={"type": "product"})).json()
    assert [e["name"] for e in ents] == ["Acme HRMS"]
    await client.patch(f"/intel/entities/{ents[0]['id']}", headers=auth(hrp), json={"confirmed": True})
    async with await tenant_session(org) as s:
        ctx = await knowledge.company_context(s)
        assert "Acme HRMS" in ctx and "Globex" not in ctx and "AI Studio" not in ctx
    r = await client.post("/intel/profile/confirm", headers=auth(hrp), json={"confirm_entities": True})
    assert r.status_code == 200 and r.json()["profile"]["confirmed"] is True
    async with await tenant_session(org) as s:
        ctx = await knowledge.company_context(s)
        assert "Globex Corporation" in ctx and "Secret Corp" not in ctx


async def test_retrieve_uses_full_text_search(client, hrp):
    org = org_of(hrp)
    async with await tenant_session(org) as s:
        rows = await knowledge.retrieve(s, "payroll compliance statutory filings", k=3)
        assert rows and set(rows[0]) == {"text", "url", "title"}
        assert "statutory filings" in rows[0]["text"]
        assert await knowledge.retrieve(s, "") == []
        # ILIKE fallback finds partial words that full-text search can't
        fallback = await knowledge.retrieve(s, "Bengalur")
        assert fallback and "Bengaluru" in fallback[0]["text"]
    r = await client.get("/intel/search", headers=auth(hrp), params={"q": "Secret Corp pilot"})
    assert r.status_code == 200 and r.json()["results"]


# ------------------------------------------------------------------ competitors
async def test_competitor_discovery_filters_and_dedupes(client, hrp):
    org = org_of(hrp)
    async with await tenant_session(org) as s:
        s.add(Competitor(organization_id=org, name="Rival Three", website="https://rival3.test/", domain="rival3.test",
                         status="rejected"))
        await s.commit()
    res = await run_inline("intel.competitors", org, {})
    assert res["suggested"] == 2
    async with await tenant_session(org) as s:
        comps = {c.domain: c for c in (await s.scalars(select(Competitor))).all()}
    assert set(comps) == {"rival1.test", "rival2.test", "rival3.test"}
    assert comps["rival3.test"].status == "rejected"  # rejected stays rejected
    assert comps["rival1.test"].status == "suggested" and comps["rival1.test"].overlap_score == 0.85
    assert comps["rival2.test"].website == "https://www.rival2.test/"
    # news, social networks, review sites, Wikipedia and our own domain never become candidates
    assert not any(d in comps for d in ("acme-hr.test", "linkedin.com", "g2.com", "wikipedia.org", "news-daily.test"))

    again = await run_inline("intel.competitors", org, {})
    assert again["suggested"] == 0
    async with await tenant_session(org) as s:
        assert len((await s.scalars(select(Competitor))).all()) == 3


async def test_competitor_crawl_detects_changes(client, hrp):
    comps = (await client.get("/intel/competitors", headers=auth(hrp), params={"status": "suggested"})).json()
    rival1 = next(c for c in comps if c["domain"] == "rival1.test")
    r = await client.patch(f"/intel/competitors/{rival1['id']}", headers=auth(hrp), json={"status": "confirmed"})
    assert r.status_code == 200 and r.json()["job_id"]
    org = org_of(hrp)
    STATE.pop("rival1_new_page", None)
    first = await run_inline("intel.competitor_crawl", org, {"competitor_id": rival1["id"]})
    assert first["pages"] == 1 and first["changed"] is False

    STATE["rival1_new_page"] = True
    RIVALS["rival1.test"] = RIVALS["rival1.test"].replace("</p>", "</p><a href='/new-feature'>New</a>")
    second = await run_inline("intel.competitor_crawl", org, {"competitor_id": rival1["id"]})
    assert second["changed"] is True
    c = next(x for x in (await client.get("/intel/competitors", headers=auth(hrp))).json() if x["domain"] == "rival1.test")
    assert c["changes"][0]["new_offerings"] == ["AI copilot"]
    assert any(u.endswith("/new-feature") for u in c["changes"][0]["new_pages"])
    assert c["content_themes"] and c["pricing_signals"] == ["per employee per month"]
    # Competitor pages never ground our own content.
    async with await tenant_session(org) as s:
        assert not [x for x in await knowledge.retrieve(s, "AI copilot for HR", k=10) if "rival1" in (x["url"] or "")]


# ------------------------------------------------------------------ audit and ideas
async def test_presence_audit(client, hrp):
    org = org_of(hrp)
    res = await run_inline("intel.audit", org, {})
    assert 0 < res["score"] <= 100
    a = (await client.get("/intel/audits/latest", headers=auth(hrp))).json()["audit"]
    sections = a["sections"]
    assert sections["technical"]["broken_link_count"] >= 1 and sections["technical"]["images_missing_alt"] == 1
    assert sections["performance"]["score"] == 62
    assert sections["content"]["latest_date"] == TODAY
    assert next(p for p in sections["social"]["platforms"] if p["platform"] == "linkedin")["linked_from_website"]
    impacts = [r["impact"] for r in a["recommendations"]]
    assert impacts == sorted(impacts, key=["high", "medium", "low"].index)


async def test_ideas_dedupe_and_pillars(client, hrp):
    org = org_of(hrp)
    async with await tenant_session(org) as s:
        pillar = ContentPillar(organization_id=org, name="Product", weight=30)
        s.add(pillar)
        await s.commit()
        pillar_id = pillar.id
    r = await client.post("/intel/opportunities", headers=auth(hrp), json={"title": "An already existing manual idea!"})
    assert r.status_code == 201
    assert (await client.post("/intel/opportunities", headers=auth(hrp), json={"title": "an already existing manual idea"})).status_code == 409

    res = await run_inline("intel.ideas", org, {})
    assert res["created"] == 3  # duplicates, empty titles and the existing idea are skipped (both rounds)
    ideas = (await client.get("/intel/opportunities", headers=auth(hrp))).json()
    titles = [knowledge.normalize_name(i["title"]) for i in ideas]
    assert len(titles) == len(set(titles))
    first = next(i for i in ideas if i["title"] == "How Acme HRMS closes payroll in hours")
    assert first["pillar_id"] == str(pillar_id) and first["formats"] == ["carousel"]
    gsc = next(i for i in ideas if i["title"] == "Search-led idea")
    assert gsc["source"] == "idea" and gsc["score"] == 0.85  # no Search Console data → not labelled gsc

    # next_opportunity picks the best idea for the pillar and marks it used
    async with await tenant_session(org) as s:
        opp = await knowledge.next_opportunity(s, pillar_id)
        await s.commit()
        assert opp.title == "How Acme HRMS closes payroll in hours" and opp.status == "used" and opp.used_count == 1
        nxt = await knowledge.next_opportunity(s, pillar_id)  # falls back to ideas without a pillar
        assert nxt is not None and nxt.pillar_id is None
        await s.rollback()

    dismissed = await client.patch(f"/intel/opportunities/{gsc['id']}", headers=auth(hrp), json={"status": "dismissed"})
    assert dismissed.json()["status"] == "dismissed"
    assert all(i["status"] == "new" for i in (await client.get("/intel/opportunities", headers=auth(hrp), params={"status": "new"})).json())


async def test_no_ai_model_fails_permanently(client, hrp):
    from app.services.jobs import PermanentError

    STATE["no_ai"] = True
    with pytest.raises(PermanentError, match="Add an AI text model"):
        await run_inline("intel.extract", org_of(hrp), {"chain": False})


async def test_run_all_chain(client, hrp):
    r = await client.post("/intel/run-all", headers=auth(hrp), json={})
    assert r.status_code == 202
    res = await run_inline("intel.run_all", org_of(hrp), {})
    assert all(step["ok"] for step in res["steps"].values()), res
    ov = (await client.get("/intel/overview", headers=auth(hrp))).json()
    assert ov["profile"]["summary"] and ov["counts"]["pages"] >= 5 and ov["latest_audit"]["score"] > 0
    assert any(j["kind"] == "intel.run_all" for j in ov["running_jobs"])  # the queued API job (no worker in tests)


# ------------------------------------------------------------------ permissions and tenancy
async def test_viewer_cannot_run_crawls(client, hrp):
    roles = {r["name"]: r["id"] for r in (await client.get("/roles", headers=auth(hrp))).json()}
    r = await client.post("/users", headers=auth(hrp),
                          json={"email": "viewer.intel@hrpai.in", "full_name": "Intel Viewer", "role_ids": [roles["Viewer"]]})
    assert r.status_code == 201, r.text
    v = await login(client, "viewer.intel@hrpai.in", r.json()["temporary_password"])
    assert (await client.get("/intel/overview", headers=auth(v))).status_code == 200
    assert (await client.get("/intel/entities", headers=auth(v))).status_code == 200
    for method, path, body in (("post", "/intel/sources/website", {"url": SITE}), ("post", "/intel/run-all", {}),
                               ("post", "/intel/competitors/discover", None), ("post", "/intel/audits/run", None),
                               ("post", "/intel/entities", {"type": "product", "name": "X"})):
        res = await getattr(client, method)(path, headers=auth(v), **({"json": body} if body is not None else {}))
        assert res.status_code == 403, path
        assert "intel.manage" in res.json()["detail"]


async def test_crawl_api_blocks_private_addresses(client, hrp):
    for url in ("http://127.0.0.1:8600", "http://localhost", "http://192.168.1.10/", "file:///etc/passwd"):
        r = await client.post("/intel/sources/website", headers=auth(hrp), json={"url": url})
        assert r.status_code == 400, url
    r = await client.post("/intel/competitors", headers=auth(hrp), json={"url": "http://10.1.2.3"})
    assert r.status_code == 400
    r = await client.post("/intel/competitors", headers=auth(hrp), json={"url": "https://acme-hr.test/pricing"})
    assert r.status_code == 400 and "own website" in r.json()["detail"]


async def test_tenant_isolation(client, hrp):
    r = await client.post("/platform/organizations", headers=auth(hrp),
                          json={"name": "Intel Other Org", "admin_email": "owner@intel-other.in", "admin_name": "Other"})
    assert r.status_code == 201, r.text
    other = await login(client, "owner@intel-other.in", r.json()["temporary_password"])
    assert (await client.get("/intel/entities", headers=auth(other))).json() == []
    assert (await client.get("/intel/competitors", headers=auth(other))).json() == []
    assert (await client.get("/intel/sources", headers=auth(other))).json() == []
    ov = (await client.get("/intel/overview", headers=auth(other))).json()
    assert ov["profile"] is None and ov["counts"]["entities_total"] == 0

    hrp_entity = (await client.get("/intel/entities", headers=auth(hrp))).json()[0]
    assert (await client.patch(f"/intel/entities/{hrp_entity['id']}", headers=auth(other), json={"name": "Hacked"})).status_code == 404
    assert (await client.delete(f"/intel/entities/{hrp_entity['id']}", headers=auth(other))).status_code == 404
    async with await tenant_session(org_of(other)) as s:
        assert await knowledge.retrieve(s, "payroll") == []
        assert await knowledge.next_opportunity(s, None) is None
        assert "Acme" not in await knowledge.company_context(s)
    async with await tenant_session(org_of(hrp)) as s:
        assert (await s.scalars(select(Opportunity))).all()
